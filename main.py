# -*- coding: utf-8 -*-
"""
Word 标题识别、格式统一与自动编号工具
依赖：pip install python-docx pywin32 openpyxl
运行：python main.py
"""

import os
import re
import shutil
import tempfile
import traceback
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    import win32com.client as win32
    import pythoncom
    HAS_COM = True
except ImportError:
    HAS_COM = False

W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


# =========================================================
# 常量
# =========================================================

TEMPLATE_PRESETS = {
    1: ["第{1}章", "第{1}篇", "第{1}部分", "第{1}编", "{1}.", "{1}、", "{1}"],
    2: ["{1}.{2}", "{1}.{2}.", "第{1}节", "（{2}）", "({2})", "{2}.", "{2}、", "{2}"],
    3: ["{1}.{2}.{3}", "{1}.{2}.{3}.", "（{3}）", "({3})", "{3}.", "{3}、", "{3}"],
    4: ["{1}.{2}.{3}.{4}", "（{4}）", "({4})", "{4}.", "{4}、", "{4}"],
    5: ["{1}.{2}.{3}.{4}.{5}", "{5}.", "{5}、", "{5}"],
}

TEMPLATE_DEFAULT = {1: "第{1}章", 2: "{1}.{2}", 3: "{1}.{2}.{3}", 4: "{4}", 5: "{5}"}

SEP_PRESETS = ["空格", "顿号（、）", "点（.）", "冒号（:）", "短横（-）", "无", "制表符"]
SEP_MAP = {
    "空格": " ", "顿号（、）": "、", "点（.）": ".", "冒号（:）": ":",
    "短横（-）": "-", "无": "", "制表符": "\t",
}

TEMPLATE_HELP_TEXT = """编号模板编写原则
============================

一、占位符
    {1}~{9}  表示第 1~9 级标题的当前序号

二、使用规则
    1. 每级模板中只能引用"本级"或"更高级"的占位符：
         一级：{1}
         二级：{1}.{2} 或 {2}
         三级：{1}.{2}.{3}、{2}.{3} 或 {3}
    2. 超出本级的占位符会被自动清除。
    3. 下级序号会随上级标题出现而自动重置为 1。
    4. 若上级标题未出现，上级序号按 1 显示（不会出现 0.1、0.0.1）。

三、常见示例
    一级模板        生成效果
    第{1}章         第1章、第2章 ...
    {1}.            1.、2.、3. ...
    二级模板        生成效果（在第 1 章下）
    {1}.{2}         1.1、1.2 ...
    第{1}节         第1节、第2节 ...
    三级模板        生成效果
    {1}.{2}.{3}     1.1.1、1.1.2 ...

四、分隔符
    编号与标题文字之间的字符，可下拉选择或手动输入。

五、注意
    {1} 生成的是阿拉伯数字（1、2、3）。
    下拉框可直接编辑，输入自定义模板后回车即可生效。
"""


# =========================================================
# 文档解析
# =========================================================

def get_outline_level(paragraph):
    pPr = paragraph._p.pPr
    if pPr is None:
        return None
    ol = pPr.find(W_NS + 'outlineLvl')
    if ol is None:
        return None
    val = ol.get(W_NS + 'val')
    if val is None:
        return None
    try:
        n = int(val) + 1
        if 1 <= n <= 9:
            return n
    except ValueError:
        pass
    return None


def get_style_level(style_name):
    if not style_name:
        return None
    m = re.match(r'^(?:Heading|标题)\s*(\d+)', style_name.strip(), re.I)
    if m:
        return int(m.group(1))
    return None


def detect_number_pattern(text):
    m = re.match(r'^(第[一二三四五六七八九十百千万零〇\d]+[章节篇部分编])\s*', text)
    if m:
        prefix = m.group(1)
        return (2 if '节' in prefix else 1), prefix

    m = re.match(r'^(\d+(?:[\.．\-]\d+)+)([\.．、\s:：\-])', text)
    if m:
        return m.group(1).replace('．', '.').count('.') + 1, m.group(0)

    m = re.match(r'^(\d+)([\.．、\s:：\)）\-])', text)
    if m:
        return 1, m.group(0)

    m = re.match(r'^([一二三四五六七八九十]+)[、.．]\s*', text)
    if m:
        return 1, m.group(0)

    m = re.match(r'^[\(（]\d+[\)）]\s*', text)
    if m:
        return 3, m.group(0)

    return None, None


def load_paragraphs(docx_path):
    doc = Document(docx_path)
    paras = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        paras.append({
            'index': i,
            'text': text,
            'style': p.style.name if p.style else '',
            'outline': get_outline_level(p),
            'is_title': False,
            'level': None,
            'source': '',
        })
    return paras


def detect_heading(para):
    if para['outline']:
        return para['outline'], '大纲级别'
    lvl = get_style_level(para['style'])
    if lvl:
        return lvl, '内置样式'
    lvl, _ = detect_number_pattern(para['text'])
    if lvl:
        return lvl, '编号模式'
    return None, None


def remove_old_number(text):
    patterns = [
        # 中文编号：第X章 / 第X节 / 第X篇 / 第X部分 / 第X编
        r'^第[一二三四五六七八九十百千万零〇\d]+[章节篇部分编][\s:：、.．\-]*',
        # 多级数字：1.1 / 1.1.1 / 1-1-1 / 1.1-2
        r'^\d+(?:[\.．\-]\d+)+[\.．、\s:：\-]*',
        # 括号数字：(1) / （1） / 【1】
        r'^[\(（\[\【]\d+[\)）\]\】][\s.．、:：\-]*',
        # 单级数字 + 分隔符：1. / 1、 / 1) / 1- / 1:
        r'^\d+[\.．、\s:：\)）\-]+\s*',
        # 纯数字后接空格：1 标题
        r'^\d+\s+',
        # 中文数字：一、二、三、
        r'^[一二三四五六七八九十]+[、.．\s:：]+',
        # 圈号：① ② ③ ...
        r'^[①-⑳]\s*',
        # 罗马数字：I. II. III. / i. ii.
        r'^(?=[IVXLCivxlc]+[\.\s])[IVXLCivxlc]+[\.\s]+',
    ]
    for pat in patterns:
        new_text = re.sub(pat, '', text, count=1)
        if new_text != text:
            text = new_text
    return text.strip()


# =========================================================
# 编号生成（★ 父级为 0 时按 1 显示）
# =========================================================

def generate_numbers(titles, templates, separators):
    """为标题序列生成编号前缀。

    - 只有 is_title=True 且 level 有效（1~9）的项才计数并生成编号。
    - 父级计数器为 0（即上级标题尚未出现）时，按 1 显示，避免 0.1、0.0.1。
    """
    counters = [0] * 10
    result = []
    for t in titles:
        if not t.get('is_title'):
            result.append('')
            continue
        level = t.get('level')
        try:
            level = int(level)
        except (TypeError, ValueError):
            result.append('')
            continue
        if level < 1 or level > 9:
            result.append('')
            continue

        counters[level - 1] += 1
        for i in range(level, 10):
            counters[i] = 0

        tmpl = templates.get(level, '{' + str(level) + '}')
        s = tmpl
        for i in range(1, level + 1):
            cnt = counters[i - 1]
            # ★ 父级计数器为 0（未出现过上级标题）时按 1 显示
            if cnt == 0:
                cnt = 1
            s = s.replace('{' + str(i) + '}', str(cnt))
        s = re.sub(r'\{\d+\}', '', s)
        result.append(s + separators.get(level, ' '))
    return result


# =========================================================
# python-docx 处理
# =========================================================

def _clear_paragraph_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def _set_paragraph_text(p, text):
    _clear_paragraph_runs(p)
    p.add_run(text)


def _replace_paragraph_text_preserving_format(p, new_text):
    runs = list(p.runs)
    if not runs:
        p.add_run(new_text)
        return
    base_run = None
    for r in runs:
        if (r.text or '').strip():
            base_run = r
            break
    if base_run is None:
        base_run = runs[0]
    base_run.text = new_text
    for r in runs:
        if r is not base_run:
            r._element.getparent().remove(r._element)


def _remove_paragraph_numbering(p):
    """移除段落上的 w:numPr，用于清除 Word 自动编号"""
    pPr = p._p.pPr
    if pPr is None:
        return
    numPr = pPr.find(qn('w:numPr'))
    if numPr is not None:
        pPr.remove(numPr)


def _remove_style_numbering(doc):
    """移除"标题 N"样式上定义的自动编号"""
    for lvl in range(1, 10):
        for name in (f"标题 {lvl}", f"Heading {lvl}"):
            try:
                st = doc.styles[name]
            except KeyError:
                continue
            el = st.element
            if el is None:
                continue
            pPr = el.find(qn('w:pPr'))
            if pPr is None:
                continue
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                pPr.remove(numPr)


def _apply_heading_style(p, doc, level):
    for name in (f"标题 {level}", f"Heading {level}"):
        try:
            p.style = doc.styles[name]
            return True
        except KeyError:
            continue
    return False


def _set_style_font_cn(style, font_name):
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)


def _apply_format_to_styles(doc, format_settings):
    if not format_settings:
        return
    for level, fmt in format_settings.items():
        for name in (f"标题 {level}", f"Heading {level}"):
            try:
                st = doc.styles[name]
            except KeyError:
                continue
            try:
                if fmt.get('font_name'):
                    _set_style_font_cn(st, fmt['font_name'])
                if fmt.get('font_size'):
                    st.font.size = Pt(fmt['font_size'])
                st.font.bold = bool(fmt.get('bold', False))
                if fmt.get('color_rgb') is not None:
                    st.font.color.rgb = RGBColor(*fmt['color_rgb'])
                if fmt.get('alignment') is not None:
                    st.paragraph_format.alignment = fmt['alignment']
                if fmt.get('space_before') is not None:
                    st.paragraph_format.space_before = Pt(fmt['space_before'])
                if fmt.get('space_after') is not None:
                    st.paragraph_format.space_after = Pt(fmt['space_after'])
            except Exception:
                traceback.print_exc()
            break


def _clean_and_style(doc, title_list, apply_heading_style=True):
    _remove_style_numbering(doc)
    for t in title_list:
        idx = t['index']
        if idx >= len(doc.paragraphs):
            continue
        p = doc.paragraphs[idx]
        clean = remove_old_number(p.text)
        _set_paragraph_text(p, clean)
        _remove_paragraph_numbering(p)
        if apply_heading_style:
            _apply_heading_style(p, doc, t['level'])


# =========================================================
# COM 自动编号
# =========================================================

def get_word_app():
    if not HAS_COM:
        return None
    try:
        pythoncom.CoInitialize()
    except Exception:
        pass
    for prog_id in ("Word.Application", "Kwps.Application", "kwps.Application"):
        try:
            app = win32.DispatchEx(prog_id)
            app.Visible = False
            app.DisplayAlerts = False
            return app
        except Exception:
            continue
    return None


def _clear_existing_list_numbers(doc, title_list):
    for t in title_list:
        try:
            para = doc.Paragraphs(t['index'] + 1)
            para.Range.ListFormat.RemoveNumbers()
        except Exception:
            pass


def apply_auto_numbering(doc, title_list, templates, separators):
    try:
        list_gallery = doc.Application.ListGalleries(2)
        try:
            list_template = list_gallery.ListTemplates(1)
        except Exception:
            list_template = list_gallery.ListTemplates.Add()

        for lvl_num in range(1, 10):
            tmpl = templates.get(lvl_num)
            if not tmpl:
                continue
            try:
                level_obj = list_template.ListLevels(lvl_num)
                w_tmpl = re.sub(r'\{(\d+)\}', lambda m: '%' + m.group(1), tmpl)
                w_tmpl = re.sub(r'\{\d+\}', '', w_tmpl)
                sep = separators.get(lvl_num, ' ')
                if sep == '\t':
                    sep_str = '\t'
                    level_obj.TrailingCharacter = 0
                else:
                    sep_str = sep
                    level_obj.TrailingCharacter = 1
                level_obj.NumberFormat = w_tmpl + sep_str
                level_obj.NumberStyle = 0
                level_obj.StartAt = 1
                level_obj.NumberPosition = 0
                level_obj.TextPosition = 0
            except Exception:
                continue

        _clear_existing_list_numbers(doc, title_list)

        for t in title_list:
            try:
                para = doc.Paragraphs(t['index'] + 1)
                for sname in (f"标题 {t['level']}", f"Heading {t['level']}"):
                    try:
                        para.Style = doc.Styles(sname)
                        break
                    except Exception:
                        continue
                try:
                    para.Range.ListFormat.ApplyListTemplateWithLevel(
                        ListTemplate=list_template,
                        ContinuePreviousList=True,
                        ApplyTo=0,
                        DefaultListBehavior=2,
                        ApplyLevel=t['level']
                    )
                except Exception:
                    try:
                        para.Range.ListFormat.ApplyListTemplate(
                            ListTemplate=list_template,
                            ContinuePreviousList=True,
                            ApplyTo=0,
                            DefaultListBehavior=2
                        )
                        try:
                            para.Range.ListFormat.ListLevelNumber = t['level']
                        except Exception:
                            pass
                    except Exception:
                        pass
            except Exception:
                continue
        return True
    except Exception:
        traceback.print_exc()
        return False


def try_com_export(src_path, out_path, title_list, format_settings,
                   templates, separators):
    if not HAS_COM:
        return False
    tmp_dir = tempfile.mkdtemp(prefix='word_title_')
    tmp_path = os.path.join(tmp_dir, 'stage1.docx')
    app = None
    try:
        doc = Document(src_path)
        _clean_and_style(doc, title_list, apply_heading_style=True)
        _apply_format_to_styles(doc, format_settings)
        doc.save(tmp_path)

        app = get_word_app()
        if app is None:
            return False
        wdoc = app.Documents.Open(os.path.abspath(tmp_path), ReadOnly=False)
        try:
            if not apply_auto_numbering(wdoc, title_list, templates, separators):
                return False
            try:
                wdoc.SaveAs2(os.path.abspath(out_path), FileFormat=16)
            except Exception:
                wdoc.SaveAs(os.path.abspath(out_path))
            return True
        finally:
            try:
                wdoc.Close(SaveChanges=False)
            except Exception:
                pass
    except Exception:
        traceback.print_exc()
        return False
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        shutil.rmtree(tmp_dir, ignore_errors=True)


# =========================================================
# 纯文本导出
# =========================================================

def text_only_export_reformat(src_path, out_path, title_list,
                              format_settings, templates, separators):
    doc = Document(src_path)
    _clean_and_style(doc, title_list, apply_heading_style=True)
    _apply_format_to_styles(doc, format_settings)

    sorted_titles = sorted(title_list, key=lambda x: x['index'])
    nums = generate_numbers(sorted_titles, templates, separators)

    for t, num in zip(sorted_titles, nums):
        idx = t['index']
        if idx >= len(doc.paragraphs):
            continue
        p = doc.paragraphs[idx]
        clean = remove_old_number(p.text)
        _set_paragraph_text(p, num + clean)
        _remove_paragraph_numbering(p)

    doc.save(out_path)


def text_only_export_renumber(src_path, out_path, title_list,
                              templates, separators):
    doc = Document(src_path)
    _remove_style_numbering(doc)

    sorted_titles = sorted(title_list, key=lambda x: x['index'])
    nums = generate_numbers(sorted_titles, templates, separators)

    target_indexes = {t['index'] for t in sorted_titles}
    for t, num in zip(sorted_titles, nums):
        idx = t['index']
        if idx not in target_indexes:
            continue
        if idx >= len(doc.paragraphs):
            continue
        p = doc.paragraphs[idx]
        clean = remove_old_number(p.text)
        _replace_paragraph_text_preserving_format(p, num + clean)
        _remove_paragraph_numbering(p)

    doc.save(out_path)


# =========================================================
# 导出主入口
# =========================================================

def export_document(src_path, out_path, titles, format_settings,
                    templates, separators,
                    use_auto_number=True, apply_format=True):
    title_list = [
        t for t in titles
        if t.get('is_title') and isinstance(t.get('level'), int)
        and 1 <= t['level'] <= 9
    ]
    title_list.sort(key=lambda x: x['index'])

    if not title_list:
        raise RuntimeError("没有可导出的有效标题")

    if not apply_format:
        text_only_export_renumber(src_path, out_path, title_list,
                                  templates, separators)
        return False

    if use_auto_number and HAS_COM:
        if try_com_export(src_path, out_path, title_list, format_settings,
                          templates, separators):
            return True

    text_only_export_reformat(src_path, out_path, title_list,
                              format_settings, templates, separators)
    return False


# =========================================================
# 导出标题清单
# =========================================================

def _export_titles_excel(titles, nums, file_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except ImportError:
        raise RuntimeError(
            "导出 Excel 需要安装 openpyxl：\n    pip install openpyxl\n"
            "或改用 .txt 保存。"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "标题清单"
    ws.append(["序号", "级别", "编号", "标题文字", "识别来源", "段落索引"])
    for i, (t, num) in enumerate(zip(titles, nums), 1):
        ws.append([i, t.get('level'), num.strip(), t.get('text', ''),
                   t.get('source', ''), t.get('index')])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center')
    for col, w in zip("ABCDEF", [6, 6, 18, 60, 12, 10]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("横向排列")
    ws2.append(["级别"] + [str(i + 1) for i in range(len(titles))])
    ws2.append(["编号"] + [num.strip() for num in nums])
    ws2.append(["标题文字"] + [t.get('text', '') for t in titles])
    for c in ws2["A"]:
        c.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 12

    wb.save(file_path)


def _export_titles_text(titles, nums, file_path):
    lines = ["序号\t级别\t编号\t标题文字\t识别来源"]
    for i, (t, num) in enumerate(zip(titles, nums), 1):
        lines.append("\t".join([
            str(i), str(t.get('level', '')), num.strip(),
            t.get('text', ''), t.get('source', ''),
        ]))
    with open(file_path, 'w', encoding='utf-8-sig') as f:
        f.write("\n".join(lines))


def export_titles_list(titles, file_path, templates, separators):
    sorted_titles = sorted(titles, key=lambda x: x['index'])
    nums = generate_numbers(sorted_titles, templates, separators)
    ext = os.path.splitext(file_path)[1].lower()
    if ext in ('.xlsx', '.xls'):
        _export_titles_excel(sorted_titles, nums, file_path)
    else:
        _export_titles_text(sorted_titles, nums, file_path)


# =========================================================
# GUI
# =========================================================

ALIGN_MAP = {
    '左对齐': WD_ALIGN_PARAGRAPH.LEFT, '居中': WD_ALIGN_PARAGRAPH.CENTER,
    '右对齐': WD_ALIGN_PARAGRAPH.RIGHT, '两端对齐': WD_ALIGN_PARAGRAPH.JUSTIFY,
}
ALIGN_NAMES = list(ALIGN_MAP.keys())


def show_template_help(parent):
    win = tk.Toplevel(parent)
    win.title("编号模板编写原则")
    win.geometry("620x560")
    txt = tk.Text(win, wrap='word', font=('Consolas', 10))
    txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
    txt.insert('1.0', TEMPLATE_HELP_TEXT)
    txt.config(state='disabled')
    ttk.Button(win, text="知道了", command=win.destroy).pack(pady=(0, 10))
    win.transient(parent)
    win.grab_set()


class ScrollableFrame(ttk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self.canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        self.vsb = ttk.Scrollbar(self, orient='vertical',
                                 command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)

        self.inner_id = self.canvas.create_window(
            (0, 0), window=self.inner, anchor='nw'
        )
        self.canvas.configure(yscrollcommand=self.vsb.set)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.vsb.pack(side=tk.RIGHT, fill=tk.Y)

        self.inner.bind('<Configure>', self._on_inner_config)
        self.canvas.bind('<Configure>', self._on_canvas_config)
        self.canvas.bind_all('<MouseWheel>', self._on_wheel)

    def _on_inner_config(self, event=None):
        self.canvas.configure(scrollregion=self.canvas.bbox('all'))

    def _on_canvas_config(self, event):
        self.canvas.itemconfig(self.inner_id, width=event.width)

    def _on_wheel(self, event):
        if not event.delta:
            return
        if abs(event.delta) >= 120:
            step = -1 * (event.delta // 120)
        else:
            step = -1 if event.delta > 0 else 1
        self.canvas.yview_scroll(step, 'units')


class LevelFormatPanel:
    def __init__(self, parent, level, defaults):
        self.level = level
        self.frame = ttk.LabelFrame(parent, text=f"{level} 级标题格式")
        self.frame.pack(fill=tk.X, padx=4, pady=2)

        row1 = ttk.Frame(self.frame)
        row1.pack(fill=tk.X, padx=4, pady=2)

        ttk.Label(row1, text="字体").pack(side=tk.LEFT)
        self.font_var = tk.StringVar(value=defaults.get('font_name', '宋体'))
        ttk.Entry(row1, textvariable=self.font_var, width=10).pack(side=tk.LEFT, padx=(2, 6))

        ttk.Label(row1, text="字号").pack(side=tk.LEFT)
        self.size_var = tk.IntVar(value=defaults.get('font_size', 14))
        ttk.Spinbox(row1, from_=8, to=72, textvariable=self.size_var,
                    width=4).pack(side=tk.LEFT, padx=(2, 6))

        self.bold_var = tk.BooleanVar(value=defaults.get('bold', True))
        ttk.Checkbutton(row1, text="加粗",
                        variable=self.bold_var).pack(side=tk.LEFT, padx=(0, 6))

        ttk.Label(row1, text="对齐").pack(side=tk.LEFT)
        self.align_var = tk.StringVar(value=defaults.get('align', '左对齐'))
        ttk.Combobox(row1, textvariable=self.align_var, values=ALIGN_NAMES,
                     width=7, state='readonly').pack(side=tk.LEFT, padx=(2, 6))

        ttk.Label(row1, text="颜色").pack(side=tk.LEFT)
        self.color_rgb = defaults.get('color_rgb', (0, 0, 0))
        self.color_btn = tk.Button(row1, text="  ", bg=self._rgb_to_hex(self.color_rgb),
                                   width=2, command=self.pick_color)
        self.color_btn.pack(side=tk.LEFT, padx=(2, 6))

        row2 = ttk.Frame(self.frame)
        row2.pack(fill=tk.X, padx=4, pady=2)
        ttk.Label(row2, text="段前(磅)").pack(side=tk.LEFT)
        self.before_var = tk.IntVar(value=defaults.get('space_before', 6))
        ttk.Spinbox(row2, from_=0, to=100, textvariable=self.before_var,
                    width=4).pack(side=tk.LEFT, padx=(2, 6))
        ttk.Label(row2, text="段后(磅)").pack(side=tk.LEFT)
        self.after_var = tk.IntVar(value=defaults.get('space_after', 6))
        ttk.Spinbox(row2, from_=0, to=100, textvariable=self.after_var,
                    width=4).pack(side=tk.LEFT, padx=(2, 6))

    @staticmethod
    def _rgb_to_hex(rgb):
        return '#%02x%02x%02x' % rgb

    def pick_color(self):
        rgb, hx = colorchooser.askcolor(color=self._rgb_to_hex(self.color_rgb))
        if rgb:
            self.color_rgb = tuple(int(x) for x in rgb)
            self.color_btn.config(bg=hx)

    def get_settings(self):
        return {
            'font_name': self.font_var.get().strip(),
            'font_size': int(self.size_var.get()),
            'bold': bool(self.bold_var.get()),
            'color_rgb': self.color_rgb,
            'alignment': ALIGN_MAP.get(self.align_var.get(), WD_ALIGN_PARAGRAPH.LEFT),
            'space_before': int(self.before_var.get()),
            'space_after': int(self.after_var.get()),
        }

    def set_enabled(self, enabled):
        state = 'normal' if enabled else 'disabled'
        for w in self.frame.winfo_children():
            try:
                w.configure(state=state)
            except Exception:
                pass
            for sub in w.winfo_children():
                try:
                    sub.configure(state=state)
                except Exception:
                    pass


class AddParagraphDialog:
    def __init__(self, parent, all_paras, existing_indexes):
        self.top = tk.Toplevel(parent)
        self.top.title("从所有段落中添加标题")
        self.top.geometry("820x480")
        self.result = []
        self.all_paras = all_paras
        self.existing = set(existing_indexes)

        ttk.Label(self.top, text="勾选要添加为标题的段落（可多选）").pack(anchor='w', padx=8, pady=6)

        cols = ("选择", "序号", "样式", "大纲", "内容")
        self.tree = ttk.Treeview(self.top, columns=cols, show='headings', height=16)
        for c in cols:
            self.tree.heading(c, text=c)
        self.tree.column("选择", width=50, anchor='center')
        self.tree.column("序号", width=60, anchor='center')
        self.tree.column("样式", width=110, anchor='w')
        self.tree.column("大纲", width=60, anchor='center')
        self.tree.column("内容", width=520, anchor='w')
        self.tree.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
        self.tree.bind('<Button-1>', self.on_click)

        self.check_state = {}
        for p in all_paras:
            iid = str(p['index'])
            mark = '✓' if p['index'] in self.existing else ''
            self.tree.insert('', tk.END, iid=iid, values=(
                mark, p['index'] + 1, p['style'],
                p.get('outline') or '', p['text'][:80]
            ))
            self.check_state[iid] = (p['index'] in self.existing)

        btns = ttk.Frame(self.top)
        btns.pack(fill=tk.X, padx=8, pady=6)
        ttk.Button(btns, text="取消", command=self.top.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btns, text="确定", command=self.confirm).pack(side=tk.RIGHT, padx=4)

        self.top.transient(parent)
        self.top.grab_set()
        parent.wait_window(self.top)

    def on_click(self, event):
        if self.tree.identify_region(event.x, event.y) != 'cell':
            return
        if self.tree.identify_column(event.x) != '#1':
            return
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        self.check_state[iid] = not self.check_state.get(iid, False)
        vals = list(self.tree.item(iid, 'values'))
        vals[0] = '✓' if self.check_state[iid] else ''
        self.tree.item(iid, values=vals)

    def confirm(self):
        self.result = [int(iid) for iid, chk in self.check_state.items() if chk]
        self.top.destroy()


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Word 标题识别、格式统一与自动编号工具")

        sw = root.winfo_screenwidth()
        sh = root.winfo_screenheight()
        w = min(1100, sw - 60)
        h = min(760, sh - 90)
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 2)
        root.geometry(f"{w}x{h}+{x}+{y}")
        root.minsize(880, 480)

        self.src_path = None
        self.all_paras = []
        self.items = []

        self.format_panels = {}
        self.template_vars = {}
        self.sep_vars = {}

        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self.root)
        top.pack(fill=tk.X, padx=8, pady=(6, 2))

        ttk.Button(top, text="选择 Word 文档",
                   command=self.select_file).pack(side=tk.LEFT)
        self.file_label = ttk.Label(top, text="未选择文件", foreground='#555')
        self.file_label.pack(side=tk.LEFT, padx=8)

        ttk.Button(top, text="识别标题",
                   command=self.detect_titles).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="从所有段落添加",
                   command=self.open_add_dialog).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="移除选中",
                   command=self.remove_selected).pack(side=tk.LEFT, padx=4)

        bottom = ttk.Frame(self.root)
        bottom.pack(fill=tk.X, padx=8, pady=(2, 8), side=tk.BOTTOM)

        ttk.Button(bottom, text="退出",
                   command=self.root.quit).pack(side=tk.RIGHT, padx=4)
        ttk.Button(bottom, text="导出为新 Word 文件",
                   command=self.export).pack(side=tk.RIGHT, padx=4)
        ttk.Button(bottom, text="导出标题清单（Excel/文本）",
                   command=self.export_titles).pack(side=tk.RIGHT, padx=4)

        scroll = ScrollableFrame(self.root)
        scroll.pack(fill=tk.BOTH, expand=True, padx=6, pady=2)
        body = scroll.inner

        mid = ttk.LabelFrame(
            body,
            text="标题列表（点击第一列勾选/取消；双击级别或标题文字列可修改）"
        )
        mid.pack(fill=tk.X, padx=2, pady=4)

        cols = ("包含", "序号", "级别", "标题文字", "来源", "段落索引")
        self.tree = ttk.Treeview(mid, columns=cols, show='headings', height=9)
        for c in cols:
            self.tree.heading(c, text=c)
        self.tree.column("包含", width=48, anchor='center')
        self.tree.column("序号", width=48, anchor='center')
        self.tree.column("级别", width=56, anchor='center')
        self.tree.column("标题文字", width=520, anchor='w')
        self.tree.column("来源", width=90, anchor='center')
        self.tree.column("段落索引", width=76, anchor='center')
        self.tree.pack(fill=tk.X, side=tk.LEFT, padx=4, pady=4)
        sb = ttk.Scrollbar(mid, orient='vertical', command=self.tree.yview)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.bind('<Button-1>', self.on_tree_click)
        self.tree.bind('<Double-1>', self.on_tree_double)

        mode_wrap = ttk.LabelFrame(body, text="处理模式")
        mode_wrap.pack(fill=tk.X, padx=2, pady=4)

        self.mode_var = tk.StringVar(value='reformat')
        ttk.Radiobutton(mode_wrap, text="重新编号 + 修改标题格式",
                        variable=self.mode_var, value='reformat',
                        command=self.on_mode_change).pack(side=tk.LEFT, padx=10, pady=2)
        ttk.Radiobutton(mode_wrap, text="只重新编号（保持原标题格式）",
                        variable=self.mode_var, value='renumber_only',
                        command=self.on_mode_change).pack(side=tk.LEFT, padx=10, pady=2)

        paned = ttk.Frame(body)
        paned.pack(fill=tk.X, padx=2, pady=2)

        left = ttk.Frame(paned)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        right = ttk.Frame(paned)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0))

        fmt_wrap = ttk.LabelFrame(left, text="按级别统一设置格式")
        fmt_wrap.pack(fill=tk.X)
        default_fonts = {1: ('黑体', 16, True), 2: ('楷体', 14, True), 3: ('宋体', 12, False)}
        for lvl in (1, 2, 3):
            fn, fs, bd = default_fonts[lvl]
            self.format_panels[lvl] = LevelFormatPanel(fmt_wrap, lvl, {
                'font_name': fn, 'font_size': fs, 'bold': bd,
                'align': '左对齐', 'color_rgb': (0, 0, 0),
                'space_before': 6 if lvl == 1 else 3,
                'space_after': 6 if lvl == 1 else 3,
            })

        num_wrap = ttk.LabelFrame(
            right,
            text="编号模板与分隔符（{1},{2},{3}… 为各级序号）"
        )
        num_wrap.pack(fill=tk.X)

        header = ttk.Frame(num_wrap)
        header.pack(fill=tk.X, padx=6, pady=(4, 2))
        ttk.Label(header, text="级别", width=6).pack(side=tk.LEFT)
        ttk.Label(header, text="编号模板").pack(side=tk.LEFT, padx=(2, 0))
        ttk.Label(header, text="分隔符").pack(side=tk.LEFT, padx=(48, 0))
        ttk.Button(header, text="📖 编写原则",
                   command=lambda: show_template_help(self.root)).pack(side=tk.RIGHT, padx=4)

        for lvl in (1, 2, 3):
            row = ttk.Frame(num_wrap)
            row.pack(fill=tk.X, padx=6, pady=3)
            ttk.Label(row, text=f"{lvl} 级", width=6).pack(side=tk.LEFT)
            tv = tk.StringVar(value=TEMPLATE_DEFAULT[lvl])
            ttk.Combobox(row, textvariable=tv,
                         values=TEMPLATE_PRESETS.get(lvl, []),
                         width=20).pack(side=tk.LEFT, padx=(2, 12))
            self.template_vars[lvl] = tv
            ttk.Label(row, text="分隔符").pack(side=tk.LEFT)
            sv = tk.StringVar(value="空格")
            ttk.Combobox(row, textvariable=sv, values=SEP_PRESETS,
                         width=9).pack(side=tk.LEFT, padx=4)
            self.sep_vars[lvl] = sv

        ttk.Label(num_wrap,
                  text="提示：下拉框可编辑。如 {1}.{2} 表示一级.二级序号。",
                  foreground='#666').pack(anchor='w', padx=8, pady=(2, 4))

        self.auto_number_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(num_wrap,
                        text="优先使用 Word 自动编号（失败自动回退纯文本）",
                        variable=self.auto_number_var).pack(anchor='w', padx=6, pady=(0, 4))

    def on_mode_change(self):
        enabled = (self.mode_var.get() == 'reformat')
        for panel in self.format_panels.values():
            panel.set_enabled(enabled)

    def select_file(self):
        path = filedialog.askopenfilename(
            title="选择 Word 文档",
            filetypes=[("Word 文档", "*.docx"), ("所有文件", "*.*")])
        if path:
            self.src_path = path
            self.file_label.config(text=path)
            self.all_paras = []
            self.items = []
            self.refresh_tree()

    def detect_titles(self):
        if not self.src_path:
            messagebox.showwarning("提示", "请先选择 Word 文档")
            return
        try:
            self.all_paras = load_paragraphs(self.src_path)
        except Exception as e:
            messagebox.showerror("错误", f"读取文档失败：{e}")
            return

        self.items = []
        for p in self.all_paras:
            level, source = detect_heading(p)
            item = dict(p)
            if level:
                item['is_title'] = True
                item['level'] = level
                item['source'] = source
            self.items.append(item)

        self.refresh_tree()
        n = sum(1 for x in self.items if x['is_title'])
        messagebox.showinfo("识别完成", f"共识别到 {n} 个标题")

    def refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        shown = [it for it in self.items if it['is_title']]
        shown.sort(key=lambda x: x['index'])
        for i, it in enumerate(shown):
            self.tree.insert('', tk.END, iid=str(i), values=(
                '✓', i + 1, it['level'], it['text'], it['source'], it['index']))

    def on_tree_click(self, event):
        if self.tree.identify_region(event.x, event.y) != 'cell':
            return
        if self.tree.identify_column(event.x) != '#1':
            return
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        vals = list(self.tree.item(iid, 'values'))
        vals[0] = '' if vals[0] == '✓' else '✓'
        self.tree.item(iid, values=vals)

    def on_tree_double(self, event):
        if self.tree.identify_region(event.x, event.y) != 'cell':
            return
        col = self.tree.identify_column(event.x)
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        if col == '#3':
            self.edit_level(iid)
        elif col == '#4':
            self.edit_text(iid)

    def edit_level(self, iid):
        vals = list(self.tree.item(iid, 'values'))
        win = tk.Toplevel(self.root)
        win.title("修改级别")
        ttk.Label(win, text="请输入 1~9 的级别：").pack(padx=10, pady=6)
        v = tk.StringVar(value=str(vals[2]))
        ent = ttk.Entry(win, textvariable=v, width=8)
        ent.pack(padx=10, pady=4)
        ent.focus_set()

        def ok():
            try:
                lv = int(v.get())
                if not 1 <= lv <= 9:
                    raise ValueError
            except ValueError:
                messagebox.showerror("错误", "级别必须是 1~9 的整数", parent=win)
                return
            vals[2] = lv
            self.tree.item(iid, values=vals)
            win.destroy()

        ttk.Button(win, text="确定", command=ok).pack(pady=6)
        win.transient(self.root)
        win.grab_set()

    def edit_text(self, iid):
        vals = list(self.tree.item(iid, 'values'))
        win = tk.Toplevel(self.root)
        win.title("修改标题文字")
        ttk.Label(win, text="编辑标题文字（可手动去掉旧编号）：").pack(padx=10, pady=6)
        txt = tk.Text(win, width=60, height=4)
        txt.pack(padx=10, pady=4)
        txt.insert('1.0', str(vals[3]))
        txt.focus_set()

        def ok():
            new_text = txt.get('1.0', 'end').strip().replace('\n', ' ')
            vals[3] = new_text
            self.tree.item(iid, values=vals)
            win.destroy()

        ttk.Button(win, text="确定", command=ok).pack(pady=6)
        win.transient(self.root)
        win.grab_set()

    def remove_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("提示", "请先在列表中选择要移除的项")
            return
        for iid in sel:
            vals = self.tree.item(iid, 'values')
            idx = int(vals[5])
            for it in self.items:
                if it['index'] == idx:
                    it['is_title'] = False
                    it['level'] = None
                    it['source'] = ''
                    break
        self.refresh_tree()

    def open_add_dialog(self):
        if not self.all_paras:
            if not self.src_path:
                messagebox.showwarning("提示", "请先选择 Word 文档")
                return
            try:
                self.all_paras = load_paragraphs(self.src_path)
            except Exception as e:
                messagebox.showerror("错误", f"读取文档失败：{e}")
                return
        if not self.items:
            self.items = [dict(p, is_title=False, level=None, source='')
                          for p in self.all_paras]

        existing = [it['index'] for it in self.items if it['is_title']]
        dlg = AddParagraphDialog(self.root, self.all_paras, existing)
        if not dlg.result:
            return

        selected = set(dlg.result)
        for it in self.items:
            it['is_title'] = it['index'] in selected
            if it['is_title'] and not it.get('level'):
                lv, src = detect_heading(it)
                it['level'] = lv or 1
                it['source'] = src or '手动标记'
            if it['is_title'] and not it.get('source'):
                it['source'] = '手动标记'
            if not it['is_title']:
                it['level'] = None
                it['source'] = ''
        self.refresh_tree()

    def _collect_templates(self):
        return {lvl: var.get().strip()
                for lvl, var in self.template_vars.items() if var.get().strip()}

    def _collect_separators(self):
        return {lvl: SEP_MAP.get(var.get(), var.get())
                for lvl, var in self.sep_vars.items()}

    def _collect_checked_titles(self):
        checked = []
        for iid in self.tree.get_children():
            vals = self.tree.item(iid, 'values')
            if vals[0] != '✓':
                continue
            try:
                idx = int(vals[5])
            except (TypeError, ValueError, IndexError):
                continue
            try:
                lv = int(vals[2])
                if not 1 <= lv <= 9:
                    raise ValueError
            except (TypeError, ValueError):
                continue
            src = ''
            for it in self.items:
                if it['index'] == idx:
                    src = it.get('source', '')
                    break
            checked.append({
                'index': idx, 'level': lv,
                'text': str(vals[3]) if len(vals) > 3 else '',
                'is_title': True, 'source': src,
            })
        return checked

    def export(self):
        if not self.src_path:
            messagebox.showwarning("提示", "请先选择 Word 文档")
            return
        if not self.items:
            messagebox.showwarning("提示", "请先识别标题")
            return

        title_list = self._collect_checked_titles()
        if not title_list:
            messagebox.showwarning("提示", "列表中没有任何被勾选、且级别有效的标题")
            return

        out_path = filedialog.asksaveasfilename(
            title="另存为新文件", defaultextension=".docx",
            filetypes=[("Word 文档", "*.docx")])
        if not out_path:
            return

        apply_format = (self.mode_var.get() == 'reformat')
        format_settings = ({lvl: panel.get_settings()
                            for lvl, panel in self.format_panels.items()}
                           if apply_format else {})
        templates = self._collect_templates()
        separators = self._collect_separators()

        used_levels = {t['level'] for t in title_list}
        missing = [lv for lv in sorted(used_levels) if lv not in templates]
        if missing:
            if not messagebox.askyesno(
                    "提示",
                    f"以下级别没有设置编号模板：{missing}\n"
                    "未设置的级别将使用默认模板 '{级别}'。\n是否继续？"):
                return

        try:
            used_auto = export_document(
                self.src_path, out_path, title_list, format_settings,
                templates, separators,
                use_auto_number=self.auto_number_var.get(),
                apply_format=apply_format)
            if apply_format:
                if used_auto:
                    messagebox.showinfo(
                        "完成",
                        f"[修改格式 + 重新编号] 使用 Word 自动编号导出：\n{out_path}")
                else:
                    messagebox.showinfo(
                        "完成",
                        f"[修改格式 + 重新编号] 使用纯文本编号导出：\n{out_path}")
            else:
                messagebox.showinfo(
                    "完成",
                    f"[只重新编号] 已使用纯文本编号导出：\n{out_path}\n\n"
                    "说明：此模式为保证原格式完全不变，不调用 Word 自动编号。")
        except Exception as e:
            traceback.print_exc()
            messagebox.showerror("错误", f"导出失败：{e}")

    def export_titles(self):
        if not self.items:
            messagebox.showwarning("提示", "请先识别标题")
            return
        title_list = self._collect_checked_titles()
        if not title_list:
            messagebox.showwarning("提示", "列表中没有任何被勾选、且级别有效的标题")
            return

        path = filedialog.asksaveasfilename(
            title="导出标题清单", defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx"), ("文本文件", "*.txt")])
        if not path:
            return

        templates = self._collect_templates()
        separators = self._collect_separators()
        try:
            export_titles_list(title_list, path, templates, separators)
            messagebox.showinfo("完成", f"已导出标题清单：\n{path}")
        except Exception as e:
            traceback.print_exc()
            messagebox.showerror("错误", f"导出失败：{e}")


def main():
    root = tk.Tk()
    try:
        style = ttk.Style()
        style.theme_use('clam')
    except Exception:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
